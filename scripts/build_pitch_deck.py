"""Build the six-slide Riskora Startup Weekend deck from local source assets.

Requires python-pptx. Screenshot assets are captured from the working app and
labelled as either observed AMC data or fictional synthetic data in the deck.
"""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
PACKAGE = ROOT / "docs" / "event-submission"
ASSETS = PACKAGE / "assets"
OUTPUT = PACKAGE / "Riskora-pitch-deck.pptx"

BG = RGBColor(247, 249, 247)
WHITE = RGBColor(255, 255, 255)
INK = RGBColor(24, 34, 31)
MUTED = RGBColor(78, 92, 86)
ACCENT = RGBColor(19, 102, 78)
PALE = RGBColor(229, 242, 235)
LINE = RGBColor(205, 216, 209)
SOFT = RGBColor(241, 245, 241)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
prs.core_properties.title = "Riskora | AIBoomi Startup Weekend Kochi"
prs.core_properties.subject = "Six-slide working MVP pitch"
prs.core_properties.author = "Ansil Muhammed N S and Muhammed Ihsan K I"
blank = prs.slide_layouts[6]


def box(slide, x, y, w, h, fill=WHITE, outline=LINE, radius=True):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    if radius:
        shape.adjustments[0] = 0.06
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    if outline is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = outline
        shape.line.width = Pt(0.8)
    return shape


def line(slide, x, y, w, color=LINE):
    box(slide, x, y, w, 0.016, fill=color, outline=None, radius=False)


def text(slide, value, x, y, w, h, size=16, color=INK, bold=False, align=None):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = Inches(0)
    frame.margin_top = frame.margin_bottom = Inches(0)
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    para = frame.paragraphs[0]
    para.text = value
    para.font.name = "Aptos Display" if size >= 26 else "Aptos"
    para.font.size = Pt(size)
    para.font.bold = bold
    para.font.color.rgb = color
    if align is not None:
        para.alignment = align
    return shape


def pill(slide, value, x, y, w, fill=PALE, color=ACCENT):
    box(slide, x, y, w, 0.36, fill=fill, outline=None)
    text(
        slide,
        value,
        x + 0.12,
        y + 0.015,
        w - 0.24,
        0.32,
        size=10,
        color=color,
        bold=True,
    )


def screenshot(slide, name, x, y, w, h):
    path = ASSETS / name
    if not path.is_file():
        raise FileNotFoundError(f"Missing captured UI asset: {path}")
    box(slide, x - 0.04, y - 0.04, w + 0.08, h + 0.08, fill=WHITE, outline=LINE)
    picture = slide.shapes.add_picture(
        str(path), Inches(x), Inches(y), width=Inches(w), height=Inches(h)
    )
    picture.name = f"Riskora UI screenshot: {name}"
    return picture


def base(number, category, title, subtitle):
    slide = prs.slides.add_slide(blank)
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = BG
    text(
        slide,
        "RISKORA  /  AIBOOMI STARTUP WEEKEND KOCHI",
        0.64,
        0.25,
        8.6,
        0.28,
        10,
        ACCENT,
        True,
    )
    text(
        slide, category.upper(), 10.2, 0.25, 2.47, 0.28, 10, MUTED, align=PP_ALIGN.RIGHT
    )
    line(slide, 0.64, 0.73, 12.05)
    text(slide, title, 0.64, 0.98, 12.0, 0.7, 30, INK, True)
    text(slide, subtitle, 0.64, 1.75, 12.0, 0.42, 14, MUTED)
    line(slide, 0.64, 6.95, 12.05)
    text(
        slide,
        f"{number:02d} / 06",
        12.0,
        7.02,
        0.68,
        0.22,
        9,
        MUTED,
        align=PP_ALIGN.RIGHT,
    )
    return slide


def notes(slide, content):
    slide.notes_slide.notes_text_frame.text = content


def row(slide, index, title, body, x, y, w):
    box(slide, x, y, w, 0.94, fill=WHITE)
    pill(slide, index, x + 0.17, y + 0.24, 0.46)
    text(slide, title, x + 0.78, y + 0.16, w - 1.02, 0.27, 15, INK, True)
    text(slide, body, x + 0.78, y + 0.48, w - 1.02, 0.31, 11, MUTED)


# Slide 1: proposition and genuine product proof.
s = base(
    1,
    "Working MVP",
    "From fund disclosures to a research shortlist.",
    "Riskora helps analysts inspect what may newly appear in a later valid portfolio statement.",
)
text(s, "Published holdings\nare retrospective.", 0.64, 2.56, 5.1, 1.06, 29, INK, True)
text(
    s,
    "We make the next research question traceable: which known stocks deserve a closer look, and what was knowable at the cutoff?",
    0.64,
    3.91,
    5.1,
    1.07,
    17,
    MUTED,
)
pill(s, "RESEARCH WORKSPACE", 0.64, 5.31, 2.08)
pill(s, "AUDITABLE RANKING", 2.87, 5.31, 2.05)
screenshot(s, "fund-detail.png", 6.17, 2.37, 6.51, 4.29)
text(
    s,
    "Actual app · ABSL Flexi Cap · observed disclosure",
    6.2,
    6.69,
    6.35,
    0.22,
    9,
    MUTED,
)
text(
    s,
    "Ansil Muhammed N S  ·  Muhammed Ihsan K I",
    0.64,
    6.61,
    5.25,
    0.25,
    10,
    ACCENT,
    True,
)
notes(
    s,
    "0:00–0:20. Riskora is a research workspace for Indian equity mutual-fund disclosures. "
    "The target is a stock newly present in a later complete, comparable statement. "
    "It does not identify an order or trade date. The screenshot is the working local app, "
    "showing an actual ABSL Flexi Cap disclosure; its publication date is unverified. "
    "Ansil worked on the frontend and interface; Ihsan worked on backend, data, and event implementation.",
)


# Slide 2: problem relevance and product insight.
s = base(
    2,
    "Problem and insight",
    "The hard part is knowing what is comparable.",
    "A missing row, an existing holding, and a late publication are different research cases.",
)
row(
    s,
    "01",
    "Reports arrive after the fact",
    "Report date and first public availability are different dates.",
    0.64,
    2.48,
    7.1,
)
row(
    s,
    "02",
    "Absence can be misleading",
    "Only a full comparable statement can establish a verified negative.",
    0.64,
    3.57,
    7.1,
)
row(
    s,
    "03",
    "Manual review is hard to audit",
    "A shortlist needs a visible cutoff, sources, and later-period checks.",
    0.64,
    4.66,
    7.1,
)
box(s, 8.04, 2.48, 4.64, 3.13, fill=ACCENT, outline=None)
text(s, "THE RISKORA RULE", 8.32, 2.79, 3.95, 0.26, 10, PALE, True)
text(s, "Unknown stays\nunknown.", 8.32, 3.2, 3.95, 0.96, 28, WHITE, True)
text(
    s,
    "Candidate inputs and labels must have been available by their cutoff. A disclosure is a position, not proof of a trade.",
    8.32,
    4.35,
    3.98,
    0.88,
    13,
    WHITE,
)
box(s, 0.64, 5.94, 12.04, 0.59, fill=SOFT, outline=None)
text(
    s,
    "Initial user: fund research analysts and small desks comparing dated Indian equity mutual-fund holdings.",
    0.91,
    6.09,
    11.48,
    0.29,
    13,
    INK,
)
notes(
    s,
    "0:20–0:50. Analysts can confuse a new position with a late publication or an incomplete "
    "statement. Riskora makes the eligibility rule explicit. If the earlier statement is "
    "incomplete, an absent stock is unknown rather than a negative example. This is the key "
    "product insight, and the target remains a later disclosed position rather than a purchase.",
)


# Slide 3: UX and end-to-end product demonstration.
s = base(
    3,
    "Product demonstration",
    "A working loop from data to candidate evidence.",
    "The live app lets a researcher inspect disclosures, evaluate a model, rank, and open the inputs.",
)
screenshot(s, "ranking-table.png", 0.64, 2.54, 8.29, 4.10)
box(s, 9.24, 2.54, 3.44, 4.10, fill=PALE, outline=None)
pill(s, "FICTIONAL SAMPLE", 9.52, 2.81, 1.92, fill=ACCENT, color=WHITE)
text(s, "01  Inspect dated holdings", 9.52, 3.38, 2.92, 0.40, 15, INK, True)
text(s, "02  Check later-period evaluation", 9.52, 4.16, 2.92, 0.50, 15, INK, True)
text(s, "03  Rank and open evidence", 9.52, 5.04, 2.92, 0.47, 15, INK, True)
text(s, "7 candidates in the local synthetic demo", 9.52, 5.87, 2.92, 0.35, 11, MUTED)
notes(
    s,
    "0:50–1:25. Show the actual application. The screenshot is from the fictional synthetic "
    "sample and is clearly labelled. Data management imports attributed CSVs; Fund explorer "
    "shows dated positions; Evaluation checks readiness and a saved holdout; Predictions ranks "
    "known candidates; Evidence shows cutoff-dated features and source links. Navigate live "
    "through these screens if the demo is available. The ranking score is relative, not a "
    "calibrated probability.",
)


# Slide 4: AI method and evaluation proof.
s = base(
    4,
    "AI and technical execution",
    "The model is testable because time is enforced.",
    "Validated inputs → cutoff-dated features → baseline comparison → a separate final holdout.",
)
row(
    s,
    "01",
    "Validate before labelling",
    "Check source URL, publication date, completeness, and ISIN.",
    0.64,
    2.49,
    5.16,
)
row(
    s,
    "02",
    "Rank with an inspectable model",
    "Logistic classifier uses prior holdings and market signals.",
    0.64,
    3.57,
    5.16,
)
row(
    s,
    "03",
    "Evaluate on later periods",
    "Walk-forward checks and final holdout versus simple baselines.",
    0.64,
    4.65,
    5.16,
)
box(s, 0.64, 5.79, 5.16, 0.72, fill=PALE, outline=None)
text(
    s,
    "Optional LLM: grounded facts from supplied excerpts; not configured in this demo.",
    0.83,
    5.92,
    4.75,
    0.48,
    11,
    ACCENT,
)
screenshot(s, "holdout-table.png", 6.12, 2.50, 6.56, 3.29)
box(s, 6.12, 5.94, 6.56, 0.57, fill=SOFT, outline=None)
text(
    s,
    "Synthetic holdout only · 14 cases, 2 additions · not real-market accuracy",
    6.32,
    6.06,
    6.19,
    0.33,
    11,
    INK,
    True,
)
notes(
    s,
    "1:25–2:05. The working AI is a scikit-learn logistic classifier. It uses features known at "
    "the cutoff, including fund preferences and optional market returns and volume. Training "
    "compares with historical-addition and fund-preference baselines on later periods, then "
    "checks a final holdout. This screenshot is the local synthetic run: only 14 holdout cases "
    "and 2 additions. It proves the evaluation code path, not Indian-market predictive power. "
    "Optional LLM event extraction is constrained to verbatim supplied excerpts and is not "
    "configured in this demo. Backend: FastAPI and SQLite; frontend: React and Vite.",
)


# Slide 5: business and user validation hypothesis.
s = base(
    5,
    "Value and viability",
    "Test an analyst workflow before setting a price.",
    "Potential value is less manual sorting and a clearer audit trail for research decisions.",
)
cards = [
    (
        0.64,
        "USER TO INTERVIEW",
        "Fund research teams",
        "Start with analysts and independent desks already comparing disclosures.",
    ),
    (
        4.71,
        "PILOT TO RUN",
        "Supervised research",
        "Use permissioned historical data; observe how teams inspect the shortlist.",
    ),
    (
        8.78,
        "MEASURE",
        "Time and trust",
        "Compare review time, evidence traceability, and ranking lift over baselines.",
    ),
]
for x, eyebrow, title, body in cards:
    box(s, x, 2.66, 3.89, 2.82, fill=WHITE)
    text(s, eyebrow, x + 0.24, 2.93, 3.38, 0.28, 10, ACCENT, True)
    text(s, title, x + 0.24, 3.42, 3.4, 0.54, 21, INK, True)
    text(s, body, x + 0.24, 4.17, 3.37, 0.88, 14, MUTED)
box(s, 0.64, 5.83, 12.03, 0.69, fill=PALE, outline=None)
text(
    s,
    "Business hypothesis: a research-team license after a successful pilot.",
    0.93,
    5.97,
    8.65,
    0.4,
    16,
    ACCENT,
    True,
)
text(
    s,
    "No customer interviews or paying users yet",
    9.28,
    6.04,
    3.1,
    0.24,
    10,
    MUTED,
    align=PP_ALIGN.RIGHT,
)
notes(
    s,
    "2:05–2:35. The initial user is a fund research analyst or small research desk. "
    "Our value hypothesis is that an auditable shortlist reduces manual sorting and helps "
    "teams explain why a candidate deserves review. We have not interviewed customers or "
    "shown willingness to pay. The next commercial test is a supervised pilot with "
    "permissioned data, measuring analyst time, evidence traceability, and ranking lift. "
    "A team license is a hypothesis, not current revenue.",
)


# Slide 6: honest status, roadmap, and specific ask.
s = base(
    6,
    "Proof and next steps",
    "A working MVP with a clear validation path.",
    "The current product demonstrates the research workflow; real-market accuracy still needs evidence.",
)
box(s, 0.64, 2.52, 5.85, 3.42, fill=WHITE)
text(s, "WORKS TODAY", 0.92, 2.82, 5.25, 0.31, 11, ACCENT, True)
text(
    s,
    "• Six-screen research app and CSV imports\n• Two AMC workbooks checked against holdings CSVs\n• Synthetic model, saved holdout, rankings, evidence",
    0.92,
    3.30,
    5.16,
    2.22,
    17,
    INK,
)
box(s, 6.77, 2.52, 5.91, 3.42, fill=PALE, outline=None)
text(s, "NEEDS PROOF", 7.06, 2.82, 5.28, 0.31, 11, ACCENT, True)
text(
    s,
    "• Verify release dates and data-use rights\n• Run real holdouts across funds and periods\n• Test usefulness with research teams",
    7.06,
    3.30,
    5.20,
    2.22,
    17,
    INK,
)
box(s, 0.64, 6.12, 12.03, 0.53, fill=ACCENT, outline=None)
text(
    s,
    "ASK: two research teams and permissioned historical data for a measured pilot",
    0.91,
    6.22,
    11.46,
    0.29,
    15,
    WHITE,
    True,
)
link = text(s, "github.com/ClashLex/Riskora", 0.64, 6.71, 4.4, 0.20, 10, ACCENT)
link.text_frame.paragraphs[0].runs[
    0
].hyperlink.address = "https://github.com/ClashLex/Riskora"
demo_link = text(
    s, "Live demo: riskora-demo.onrender.com", 5.02, 6.71, 5.25, 0.20, 10, ACCENT
)
demo_link.text_frame.paragraphs[0].runs[
    0
].hyperlink.address = "https://riskora-demo.onrender.com/"
notes(
    s,
    "2:35–3:00. We have a working product and have checked the two supplied AMC workbooks "
    "against their normalized CSVs. Their exact public release dates are not verified, so "
    "Riskora labels them observed and excludes them from model training. The trained demo "
    "uses fictional data only. Next we need publication-verified history, real holdouts, "
    "and analyst pilots. Our ask is two research teams and permissioned historical data.\n\n"
    "Q&A: Does it predict trades? No: the target is a new later disclosed position. "
    "Is the accuracy real? No: the saved run is synthetic. Why AI? The classifier orders "
    "candidates from dated signals; optional LLM extraction is separate. Who pays? A team "
    "license is a testable hypothesis. Event history: the owner says the final Riskora "
    "version was built on 24 October after several idea changes; Git also records earlier "
    "exploratory work. No prior-prototype declaration or assigned surprise element was "
    "reported, and the official build-window times remain unconfirmed.\n\n"
    "Source notes: code and data review in this repository; local UI captures on 24 Oct 2026; "
    "handbook timing and six-slide limit summarized in docs/event-submission/README.md. "
    "The live Render demo is linked on this slide and contains fictional synthetic data only. "
    "Keep the separate AI impact statement for submission.",
)

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
prs.save(OUTPUT)
print(f"Wrote {len(prs.slides)} slides to {OUTPUT}")
